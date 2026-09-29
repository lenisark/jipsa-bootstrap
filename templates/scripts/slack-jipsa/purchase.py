"""비품 구매기록·월간분석 순수 로직. supply 모듈의 파싱/배치매칭 재활용."""
from __future__ import annotations
import json, re, importlib.util
from pathlib import Path

def _norm(s: str) -> str:
    s = re.sub(r'[\s()（）]+', '', (s or ''))
    return s.lower()

def normalize_dept(raw: str, known: list[str], aliases: dict) -> tuple[str, bool]:
    raw = (raw or '').strip()
    if not raw:
        return '', True
    if raw in known:
        return raw, False
    if raw in aliases and aliases[raw] in known:
        return aliases[raw], False
    nr = _norm(raw)
    for k in known:                       # 정규화 부분일치(짧은 쪽이 긴 쪽에 포함)
        nk = _norm(k)
        if nr and (nr in nk or nk in nr):
            return k, False
    return raw, True

CATEGORIES = ['사무용품','다과·음료','청소·위생','IT·전자','비품·가구','기타']
USES = ['사내비품','공용(사내·외부 혼재)','재판매·대고객']

def classify_items(items: list[str], run_claude, guide_text: str = '') -> dict:
    items = [i for i in dict.fromkeys(items) if i]     # 중복 제거·순서 보존
    if not items:
        return {}
    prompt = (
        "다음 비품 품목들을 각각 '용도'와 '카테고리'로 분류해 JSON 배열로만 출력하라.\n"
        f"카테고리(택1): {', '.join(CATEGORIES)}\n"
        f"용도(택1): {', '.join(USES)}. 기본값은 '사내비품'.\n"
        + (f"분류 가이드:\n{guide_text}\n" if guide_text else "")
        + '각 원소는 {"품목":문자열,"용도":문자열,"카테고리":문자열}. JSON 외 출력 금지.\n\n'
        + '품목:\n' + '\n'.join(f'- {i}' for i in items))
    out = (run_claude(prompt) or '').strip()
    m = re.search(r'\[.*\]', out, re.S)
    if not m:
        return {}
    try:
        data = json.loads(m.group(0))
    except Exception:
        return {}
    res = {}
    for d in data:
        if not isinstance(d, dict):
            continue
        name = (d.get('품목') or '').strip()
        cat = (d.get('카테고리') or '').strip()
        use = (d.get('용도') or '').strip()
        if not name:
            continue
        res[name] = {'용도': use if use in USES else '사내비품',
                     '카테고리': cat if cat in CATEGORIES else '기타'}
    return res

def classify_with_cache(items, cache: dict, run_claude, guide_text: str = ''):
    cache = dict(cache)
    misses = [i for i in dict.fromkeys(items) if i and i not in cache]
    if misses:
        fresh = classify_items(misses, run_claude, guide_text)
        cache.update(fresh)
    result = {i: cache[i] for i in dict.fromkeys(items) if i in cache}
    return result, cache

REC_HEADERS = ['일자','부서','용도','카테고리','품목','수량','단가','금액']

def decide_use(llm_use: str, dept: str, shared_depts) -> str:
    """분류 가이드의 용도 규칙: 재판매·대고객은 품목으로(LLM 판단 유지), 공용은 발주 부서로
    (shared_depts, 예: 매입부·제휴매입파트), 나머지는 사내비품. shared_depts=None 이면 LLM 판단 그대로."""
    if shared_depts is None or llm_use == '재판매·대고객':
        return llm_use
    return '공용(사내·외부 혼재)' if dept in shared_depts else '사내비품'


def build_records(rows, classify_map, when_ymd, known_depts, dept_aliases, shared_depts=None):
    recs, warns = [], []
    for r in rows:
        name = (r.get('품명') or '').strip()
        if not name:
            continue
        qty = int(r.get('수량') or 0)
        amt = int(r.get('금액') or 0)
        dept, is_new = normalize_dept(r.get('부서',''), known_depts, dept_aliases)
        if is_new and dept:
            warns.append(f'신규 부서 후보: "{dept}" (마스터에 없음)')
        cls = classify_map.get(name, {'용도':'사내비품','카테고리':'기타'})
        recs.append({'일자': r.get('날짜') or when_ymd, '부서': dept,
                     '용도': decide_use(cls['용도'], dept, shared_depts),
                     '카테고리': cls['카테고리'], '품목': name, '수량': qty,
                     '단가': round(amt/qty) if qty > 0 else 0, '금액': amt})
    return recs, warns

def _amt(v):
    try: return int(float(str(v).replace(',', '')))
    except (TypeError, ValueError): return 0

def _day_of(v):
    s = str(v or '')
    m = re.search(r'(\d{1,2})\s*일', s)
    if m: return f'{int(m.group(1))}일'
    m = re.search(r'\d{4}-\d{2}-(\d{2})', s)
    if m: return f'{int(m.group(1))}일'
    return s

def log_rows_to_integrated(log_rows, aliases=None):
    """누적 구매로그 행 → 통합원본 행. 로그의 블록ID 보존, 일자는 'N일'로 변환.
    aliases(dept_aliases)는 병합 시점에도 적용 — 별칭 추가 전에 로그에 남은 표기(예: 손님용)를 흡수."""
    aliases = aliases or {}
    out=[]
    for r in log_rows:
        dept = str(r.get('부서','') or '').strip()
        out.append({'월':str(r.get('월','')).strip(),'일자':_day_of(r.get('일자')),
                    '부서':aliases.get(dept, dept),'용도':r.get('용도',''),
                    '카테고리':r.get('카테고리',''),'품목':r.get('품목',''),
                    '금액':_amt(r.get('금액')),'블록ID':r.get('블록ID','')})
    return out

def compute_pivots(rows) -> dict:
    from collections import defaultdict
    dm = defaultdict(int); cm = defaultdict(int); um = defaultdict(int)
    du = defaultdict(int); dc = defaultdict(int); mo = defaultdict(int)
    items = []; total = 0
    for r in rows:
        a = _amt(r.get('금액'))
        dep, mon = r.get('부서',''), r.get('월','')
        use, cat = r.get('용도',''), r.get('카테고리','')
        dm[(dep,mon)] += a; cm[(cat,mon)] += a; um[(use,mon)] += a
        du[(dep,use)] += a; dc[(dep,cat)] += a; mo[mon] += a; total += a
        # top20 튜플: (금액, 월, 부서, 용도, 카테고리, 품목) — 큰지출_TOP20 시트 열 순서
        items.append((a, mon, dep, use, cat, r.get('품목','')))
    items.sort(key=lambda x: x[0], reverse=True)
    return {'부서월':dict(dm),'카테고리월':dict(cm),'용도월':dict(um),
            '부서용도':dict(du),'부서카테고리':dict(dc),'월합':dict(mo),
            'top20':items[:20],'총계':total}

def _sibling(mod_file):
    spec = importlib.util.spec_from_file_location(mod_file[:-3], Path(__file__).with_name(mod_file))
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m

def load_guide_text(cfg, max_chars: int = 3000) -> str:
    """최신 분석 파일의 '분류_가이드' 시트를 줄글로(분류 프롬프트에 주입). 없거나 못 읽으면 ''."""
    try:
        import openpyxl
        pstore = _sibling('purchase_store.py')
        path, _ = pstore.latest_analysis(cfg['folder'], cfg.get('analysis_prefix', ''))
        if not path:
            return ''
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if '분류_가이드' not in wb.sheetnames:
            wb.close()
            return ''
        lines = [' | '.join(str(c).strip() for c in r if c not in (None, ''))
                 for r in wb['분류_가이드'].iter_rows(values_only=True)]
        wb.close()
        return '\n'.join(ln for ln in lines if ln)[:max_chars]
    except Exception:
        return ''


def apply_purchase_record(cfg, rows, run_claude, when_ymd) -> dict:
    """붙여넣기 행들 → 분류 → 데몬 전용 누적 구매로그에 append(사람 양식 미접촉).
    반환 {'appended','records'(로그행),'warns','error'}."""
    pstore = _sibling('purchase_store.py')
    log_path = Path(cfg['folder']) / cfg.get('purchase_log', '비품 구매기록_자동.xlsx')
    if pstore.is_locked(log_path):
        return {'appended': 0, 'records': [], 'warns': [], 'error': 'locked'}
    cache_path = Path(cfg['classify_cache']).expanduser()
    try:
        cache = json.loads(cache_path.read_text(encoding='utf-8')) if cache_path.exists() else {}
    except Exception:
        cache = {}
    names = [(r.get('품명') or '').strip() for r in rows if (r.get('품명') or '').strip()]
    # 분류_가이드 시트 주입은 guide_sheet: true 일 때만(가이드 문구와 팀 관행이 다를 수 있어 opt-in)
    guide = cfg.get('guide_text') or (load_guide_text(cfg) if cfg.get('guide_sheet') else '')
    cmap, cache2 = classify_with_cache(names, cache, run_claude, guide)
    recs, warns = build_records(rows, cmap, when_ymd,
                                cfg.get('known_depts', []), cfg.get('dept_aliases', {}),
                                cfg.get('shared_use_depts'))
    # 월·블록ID는 행 일자 기준(날짜 칸 없으면 붙여넣은 날). seq는 그 달 기존 로그 건수에서 이어짐.
    existing = pstore.read_purchase_log(log_path)
    seqs: dict[str, int] = {}
    log_rows = []
    for rec in recs:
        yyyymm = rec['일자'][:7].replace('-', '')
        month_label = f'{int(yyyymm[4:6])}월'
        if yyyymm not in seqs:
            seqs[yyyymm] = sum(1 for r in existing
                               if str(r.get('블록ID', '')).startswith(f'{yyyymm}#')) + 1
        log_rows.append({'월': month_label, '일자': rec['일자'], '부서': rec['부서'],
                         '용도': rec['용도'], '카테고리': rec['카테고리'], '품목': rec['품목'],
                         '수량': rec['수량'], '단가': rec['단가'], '금액': rec['금액'],
                         '블록ID': f'{yyyymm}#{seqs[yyyymm]}'})
        seqs[yyyymm] += 1
    n = pstore.append_purchase_log(log_path, log_rows)
    if cache2 != cache:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(json.dumps(cache2, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'appended': n, 'records': log_rows, 'warns': warns, 'error': None}


def _yymmdd(ymd):
    return ymd[2:].replace('-', '')          # 2026-08-06 -> 260806


def merge_month_into_analysis(cfg, yyyymm, when_ymd, dry_run=True) -> dict:
    """최신 분석 로드 → 구매로그에서 해당 월 중 미반영(블록ID) 행 읽기 → 통합원본 변환 →
    dry_run이면 제안만, 아니면 새 버전 파일 저장(원본 불변). 결정론적 core(LLM 미사용).
    반환 {'status':'proposed'|'merged'|'nothing'|'locked','month','rows','out','summary'}."""
    pstore = _sibling('purchase_store.py')
    folder = Path(cfg['folder'])
    month_label = f'{int(yyyymm[4:6])}월'
    nothing = {'status': 'nothing', 'month': month_label, 'rows': 0, 'out': None, 'summary': {}}
    analysis, ver = pstore.latest_analysis(folder, cfg['analysis_prefix'])
    if not analysis:
        return nothing
    integ = pstore.read_integrated(analysis)
    log_path = folder / cfg.get('purchase_log', '비품 구매기록_자동.xlsx')
    if not log_path.exists():
        return nothing                       # 구매로그 없음
    if pstore.is_locked(analysis) or pstore.is_locked(log_path):
        return {'status': 'locked', 'month': month_label, 'rows': 0, 'out': None, 'summary': {}}
    # 블록ID 단위 증분: 같은 달을 이미 병합한 뒤 로그에 추가된 행만 반영(월 단위 판정은 후행 입고를 영구 누락시켰음)
    have = {str(r.get('블록ID', '')).strip() for r in integ}
    month_rows = [r for r in pstore.read_purchase_log(log_path)
                  if str(r.get('월', '')).strip() == month_label
                  and str(r.get('블록ID', '')).strip() not in have]
    new_rows = log_rows_to_integrated(month_rows, cfg.get('dept_aliases', {}))
    if not new_rows:
        return nothing                       # 그 달 로그 없음 또는 전부 이미 반영
    all_rows = integ + new_rows
    pivots = compute_pivots(all_rows)
    summary = {'월합': pivots['월합'], '총계': pivots['총계'], '건수': len(new_rows)}
    if dry_run:
        return {'status': 'proposed', 'month': month_label, 'rows': len(new_rows),
                'out': None, 'summary': summary}
    nv = pstore.next_version(ver)
    out = folder / f"{_yymmdd(when_ymd)}-{cfg['dept_code']}-{cfg['analysis_prefix']}-v{nv[0]}.{nv[1]}.xlsx"
    pstore.write_merged_analysis(analysis, out, new_rows, pivots)
    return {'status': 'merged', 'month': month_label, 'rows': len(new_rows),
            'out': str(out), 'summary': summary}


# ── 주문 캡처 → 구매행 (2026-09) ────────────────────────────────────
CAPTURE_PROMPT = (
    "다음 이미지들은 쿠팡 등 온라인몰 주문목록 캡처다. 각 이미지를 Read 도구로 열어 구매 품목을 "
    "JSON 배열로만 출력하라. 각 원소는 "
    '{"품명": 상품명(핵심 옵션 포함, 40자 이내), "수량": 정수, "단가": 정수(원), "날짜": "YYYY-MM-DD"}.\n'
    "규칙:\n"
    "- 쿠팡의 'N,NNN 원 · K개'는 1개당 가격(단가)과 수량이다.\n"
    "- '분리 배송'·'분리배송'·'일부 상품이 분리되어 배송됩니다'·'분리배송된 상품입니다' 표시가 있는 카드는 "
    "같은 상품 하나를 나눠 보낸 것이다. 같은 상품의 분리배송 카드가 여러 장이어도 한 줄만 적고, "
    "수량은 카드 한 장에 적힌 수량 그대로 쓴다(카드 장수만큼 더하지 말 것).\n"
    "- 'K개 중 1개'라고 적힌 카드는 한 상품(수량 K)을 나눠 배송한 것이다. 같은 상품의 그런 카드가 여러 장이면 "
    "한 줄만, 수량 K로 적는다.\n"
    "- 그 밖의 카드는 각각 적는다. 이미지끼리 겹쳐 같은 카드가 두 번 보이면 한 번만 적는다.\n"
    "- 품명 앞의 배송 배지(로켓, 판매자로켓, 로켓프레시, 로켓직구 등)는 품명에서 뺀다.\n"
    "- 날짜는 '2026. 9. 17 주문' 같은 주문일 헤더를 따른다. 헤더가 안 보이는 이미지는 스크롤을 이어 찍은 것이니 "
    "바로 앞 이미지의 마지막 주문일을 이어받는다. 첫 이미지부터 헤더가 없으면 빈 문자열(도착일로 추정하지 말 것).\n"
    "- 취소·반품된 카드는 뺀다. JSON 외에는 아무것도 출력하지 마라.\n\n이미지 파일:\n")


def parse_capture_json(out: str, dept: str = '') -> list[dict]:
    """캡처 판독 결과(JSON) → 구매표 행 [{품명,수량,금액,부서,날짜}]. 금액 = 단가×수량."""
    m = re.search(r'\[.*\]', out or '', re.S)
    if not m:
        return []
    try:
        data = json.loads(m.group(0))
    except Exception:
        return []
    rows = []
    for d in data if isinstance(data, list) else []:
        if not isinstance(d, dict) or not str(d.get('품명', '')).strip():
            continue
        try:
            qty = int(d.get('수량') or 0)
            unit = _amt(d.get('단가'))
        except (TypeError, ValueError):
            continue
        if qty <= 0 or unit <= 0:
            continue
        day = str(d.get('날짜') or '').strip()
        rows.append({'품명': str(d['품명']).strip()[:60], '수량': qty, '금액': unit * qty,
                     '부서': dept, '날짜': day if re.fullmatch(r'\d{4}-\d{2}-\d{2}', day) else ''})
    return rows


def rows_to_table(rows: list[dict]) -> str:
    """구매표 행 → `입고등록`에 그대로 다시 붙여넣을 수 있는 | 표."""
    lines = ['품명 | 수량 | 금액 | 부서 | 날짜']
    for r in rows:
        lines.append(f"{r['품명']} | {r['수량']} | {r['금액']:,} | {r.get('부서') or ''} | {r.get('날짜') or ''}")
    return '\n'.join(lines)
